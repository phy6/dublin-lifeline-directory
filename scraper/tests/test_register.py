import json
import os
import sys

import pytest

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))
sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", ".."))

from scraper.fetch_register import filter_candidates, is_relevant, parse_register, run

try:
    import openpyxl  # noqa: F401

    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False


def _row(**kw):
    base = {
        "rcn": "20000001",
        "name": "Test Org",
        "aka": None,
        "status": "Registered",
        "classification": "",
        "address": "1 Test Street, Dublin 8, Dublin, Republic of Ireland",
        "purpose": "Relief of poverty or economic hardship",
    }
    base.update(kw)
    return base


def test_relevant_registered_dublin_poverty():
    assert is_relevant(_row()) is True


def test_rejects_deregistered():
    assert is_relevant(_row(status="Deregistered")) is False


def test_rejects_non_dublin():
    assert is_relevant(_row(address="4 Lapp's Quay, Cork, Republic of Ireland")) is False


def test_rejects_irrelevant_purpose():
    assert is_relevant(_row(purpose="Advancement of religion", classification="")) is False


def test_accepts_homeless_classification():
    row = _row(
        classification="Social and community services [Homeless services]",
        purpose="Something else entirely",
    )
    assert is_relevant(row) is True


def test_accepts_community_welfare_purpose():
    row = _row(purpose="Advancement of community welfare including relief")
    assert is_relevant(row) is True


def test_filter_candidates_keeps_only_relevant():
    rows = [
        _row(rcn="1"),
        _row(rcn="2", status="Deregistered"),
        _row(rcn="3", address="Galway"),
    ]
    assert [r["rcn"] for r in filter_candidates(rows)] == ["1"]


def _write_workbook(path):
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Public Register"
    ws["A1"] = "Effective Date "
    ws["B1"] = "2026-09-09"
    headers = [
        "Registered Charity Number", "Registered Charity Name", "Also Known As",
        "Status", "Charity Classification: Primary", "Primary Address",
        "Also Operates In", "Governing Form", "CRO Number", "Country Established",
        "Charitable Purpose", "Charitable Objects", "Trustees (Start Date)",
    ]
    for i, h in enumerate(headers, start=1):
        ws.cell(row=2, column=i, value=h)
    data = [
        ["20000001", "Dublin Test Charity", None, "Registered", "",
         "1 Test Street, Dublin 8", None, "CLG", None, "Ireland",
         "Relief of poverty or economic hardship", None, None],
        ["20000002", "Cork Test Charity", None, "Registered", "",
         "4 Quay Street, Cork", None, "CLG", None, "Ireland",
         "Relief of poverty or economic hardship", None, None],
    ]
    for ri, row in enumerate(data, start=3):
        for ci, v in enumerate(row, start=1):
            ws.cell(row=ri, column=ci, value=v)
    wb.save(path)


def test_parse_and_run_end_to_end(tmp_path):
    if not HAS_OPENPYXL:
        pytest.skip("openpyxl not installed")
    xlsx = str(tmp_path / "register.xlsx")
    _write_workbook(xlsx)
    rows, meta = parse_register(xlsx)
    assert len(rows) == 2
    assert meta.get("effective_date") == "2026-09-09"
    assert rows[0]["rcn"] == "20000001"
    out = str(tmp_path / "candidates.json")
    result = run(xlsx, out)
    assert result["counts"] == {
        "total": 2, "registered": 2, "dublin_registered": 1, "candidates": 1,
    }
    assert result["candidates"][0]["name"] == "Dublin Test Charity"
    with open(out) as f:
        assert json.load(f)["counts"]["candidates"] == 1
