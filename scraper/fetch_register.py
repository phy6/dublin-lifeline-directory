"""Download and filter the Charities Regulator Public Register (XLSX).

The regulator's HTML pages sit behind a Cloudflare challenge, but the
published register spreadsheet downloads directly. This script fetches it,
keeps charities that are Registered + Dublin-based + relevant to
homelessness/poverty relief, and writes review candidates as JSON.

Usage:
    python3 scraper/fetch_register.py --download
    python3 scraper/fetch_register.py --xlsx /tmp/register.xlsx
"""

import argparse
import json
import logging
import os
from datetime import datetime, timezone

import httpx

logger = logging.getLogger(__name__)

REGISTER_XLSX_URL = (
    "https://www.charitiesregulator.ie/media/5rrnldzg/public-register-of-charities.xlsx"
)
HEADER_ROW = 2
DATA_START_ROW = 3
MAX_CONSECUTIVE_EMPTY = 100


def parse_register(path: str) -> tuple[list[dict], dict]:
    """Parse the register workbook. Returns (rows, meta)."""
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb["Public Register"]
    meta: dict = {}
    try:
        top_left = ws["A1"].value
        top_next = ws["B1"].value
        if top_left and "effective" in str(top_left).lower():
            meta["effective_date"] = str(top_next)
    except Exception:
        pass
    rows: list[dict] = []
    empty = 0
    for r in ws.iter_rows(min_row=DATA_START_ROW, values_only=True):
        if not r or r[0] is None:
            empty += 1
            if empty >= MAX_CONSECUTIVE_EMPTY:
                break
            continue
        empty = 0
        rows.append(
            {
                "rcn": str(r[0]),
                "name": str(r[1] or "").strip(),
                "aka": str(r[2]).strip() if r[2] else None,
                "status": str(r[3] or "").strip(),
                "classification": str(r[4] or "").strip(),
                "address": str(r[5] or "").strip(),
                "purpose": str(r[10] or "").strip() if len(r) > 10 else "",
            }
        )
    wb.close()
    return rows, meta


def is_relevant(row: dict) -> bool:
    """Registered + Dublin address + homelessness/poverty/community-welfare remit."""
    if row["status"] != "Registered":
        return False
    if "dublin" not in row["address"].lower():
        return False
    classification = row["classification"].lower()
    purpose = row["purpose"].lower()
    return (
        "homeless" in classification
        or "poverty" in purpose
        or "community welfare" in purpose
    )


def filter_candidates(rows: list[dict]) -> list[dict]:
    return [r for r in rows if is_relevant(r)]


def download_register(url: str, dest: str, timeout: float = 120.0) -> str:
    headers = {"User-Agent": "Mozilla/5.0 (X11; Linux x86_64) AppleWebKit/537.36"}
    with httpx.stream("GET", url, headers=headers, follow_redirects=True, timeout=timeout) as r:
        r.raise_for_status()
        os.makedirs(os.path.dirname(os.path.abspath(dest)), exist_ok=True)
        with open(dest, "wb") as f:
            for chunk in r.iter_bytes():
                f.write(chunk)
    return dest


def run(xlsx_path: str, output_path: str) -> dict:
    rows, meta = parse_register(xlsx_path)
    registered = [r for r in rows if r["status"] == "Registered"]
    dublin = [r for r in registered if "dublin" in r["address"].lower()]
    candidates = filter_candidates(rows)
    result = {
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "source_url": REGISTER_XLSX_URL,
        "effective_date": meta.get("effective_date"),
        "counts": {
            "total": len(rows),
            "registered": len(registered),
            "dublin_registered": len(dublin),
            "candidates": len(candidates),
        },
        "candidates": candidates,
    }
    os.makedirs(os.path.dirname(os.path.abspath(output_path)), exist_ok=True)
    with open(output_path, "w") as f:
        json.dump(result, f, indent=2)
    logger.info(
        "Register: %d total, %d registered, %d Dublin, %d candidates -> %s",
        len(rows), len(registered), len(dublin), len(candidates), output_path,
    )
    return result


def parse_args(args=None):
    parser = argparse.ArgumentParser(description="Fetch and filter the Charities Register")
    parser.add_argument("--xlsx", default=None, help="Path to a downloaded register .xlsx")
    parser.add_argument("--download", action="store_true", help="Download the register first")
    parser.add_argument("--url", default=REGISTER_XLSX_URL, help="Register XLSX URL override")
    parser.add_argument("--output", default=None, help="Candidates JSON output path")
    parser.add_argument("--verbose", action="store_true")
    return parser.parse_args(args)


def main(args=None) -> int:
    parsed = parse_args(args)
    logging.basicConfig(level=logging.DEBUG if parsed.verbose else logging.INFO)
    project_root = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    default_dir = os.path.join(project_root, "scraper", "output")
    if parsed.download or not parsed.xlsx:
        dest = parsed.xlsx or os.path.join(default_dir, "public-register-of-charities.xlsx")
        logger.info("Downloading register to %s", dest)
        download_register(parsed.url, dest)
        xlsx_path = dest
    else:
        xlsx_path = parsed.xlsx
    output_path = parsed.output or os.path.join(default_dir, "register_candidates.json")
    run(xlsx_path, output_path)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
